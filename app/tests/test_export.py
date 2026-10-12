"""Excel export: the file opens cleanly and its cells are usable - phone
numbers keep their "+", the score is a number, the time a real date, links
are clickable, empty values are empty cells, and the formula guard still
neutralizes injected formulas.
Run:  python tests/test_export.py  (from the app/ directory)"""

from __future__ import annotations

import os
import sys
from _scratch import scratch_dir  # noqa: E402  (tests/_scratch.py)
import zipfile
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from backend.collector.exporter import COLUMNS, write_category_file  # noqa: E402


def main():
    out = scratch_dir(prefix="export_test_")
    records = [
        {"Company Name": "Acme Tax LLP", "Category": "Finance",
         "Official Website": "https://acme.example", "Business Email": "info@acme.example",
         "Business Phone": "+91 98765 43210", "Country": "India",
         "Source URL": "https://acme.example/contact", "Collected At": "2026-10-12T21:58:48",
         "Confidence Score": "85"},
        {"Company Name": "=HYPERLINK(\"http://evil\",\"x\")", "Category": "Finance",
         "Business Phone": "+SUM(1,1)", "Services": "@cmd", "Business Email": "not an email",
         "Official Website": "javascript:alert(1)", "Confidence Score": "",
         "Collected At": "yesterday"},
    ]
    write_category_file("Finance", "01_Finance.xlsx", records, out)
    path = os.path.join(out, "01_Finance.xlsx")
    assert zipfile.ZipFile(path).testzip() is None
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS
    col = {name: i for i, name in enumerate(COLUMNS)}
    good, evil = list(ws[2]), list(ws[3])

    assert good[col["Business Phone"]].value == "+91 98765 43210"
    print("  phone numbers keep their leading +")
    assert good[col["Confidence Score"]].value == 85
    print("  confidence score is a number")
    when = good[col["Collected At"]]
    assert when.value == datetime(2026, 10, 12, 21, 58, 48) and "yy" in when.number_format
    print("  collected-at is a real date")
    assert good[col["Official Website"]].hyperlink.target == "https://acme.example"
    assert good[col["Business Email"]].hyperlink.target == "mailto:info@acme.example"
    print("  website / email are clickable")
    assert good[col["Subcategory"]].value is None
    print("  missing values are empty cells")
    assert ws.auto_filter.ref == "A1:Q3" and ws.freeze_panes == "A2"
    print("  header has filter arrows and stays frozen")

    for name in ("Company Name", "Business Phone", "Services"):
        v = evil[col[name]].value
        assert isinstance(v, str) and not v.startswith(("=", "+", "@")), (name, v)
        assert evil[col[name]].data_type != "f"
    assert evil[col["Official Website"]].hyperlink is None
    assert evil[col["Business Email"]].hyperlink is None
    assert evil[col["Collected At"]].value == "yesterday"
    print("  injected formulas neutralized; no links to non-web targets")
    print("ALL EXPORT TESTS PASSED")


if __name__ == "__main__":
    main()
